
import json
from pathlib import Path
from docx import Document

from loss_registry import LOSSY_TYPES

ROOT = Path(__file__).resolve().parent
INPUT = ROOT / "source.docx"
OUTPUT_MD = ROOT / "document.md"
OUTPUT_LOSS = ROOT / "loss_report.json"

def convert():
    doc = Document(str(INPUT))
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
    for table in doc.tables:
        lines.append("")
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    OUTPUT_MD.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    manifest = json.loads((ROOT / "source_features.json").read_text(encoding="utf-8"))
    features = manifest.get("unrepresentable_or_partially_representable", [])
    # BUG: starter only reports one lossy category even though the feature inventory contains several.
    loss = [
        {"element_type": feature["element_type"], "description": feature["description"]}
        for feature in features
        if feature["element_type"] in LOSSY_TYPES
    ][:1]
    OUTPUT_LOSS.write_text(json.dumps(loss, indent=2), encoding="utf-8")

if __name__ == "__main__":
    convert()
