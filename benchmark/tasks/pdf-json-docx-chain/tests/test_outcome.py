
import json
import re
import subprocess
import sys
import unittest
from decimal import Decimal
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
TEST_DIR = Path(__file__).resolve().parent
PROJECT = ROOT / "project"
SCRIPT = PROJECT / "pipeline.py"
EXPECTED = json.loads((TEST_DIR / "expected_financials.json").read_text(encoding="utf-8"))
JSON_OUT = PROJECT / "extracted.json"
DOCX_OUT = PROJECT / "report.docx"

def run_pipeline():
    for p in [JSON_OUT, DOCX_OUT]:
        if p.exists():
            p.unlink()
    subprocess.run([sys.executable, str(SCRIPT)], check=True, cwd=str(PROJECT))
    data = json.loads(JSON_OUT.read_text(encoding="utf-8"))
    doc = Document(str(DOCX_OUT))
    return data, doc

def doc_all_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)

def numeric_values(text):
    """Parse human-formatted numbers such as $125,000 and 125000 equally."""
    return {
        Decimal(token.replace(",", ""))
        for token in re.findall(r"(?<![A-Za-z0-9])[-+]?\d[\d,]*(?:\.\d+)?", text)
    }

class HiddenTests(unittest.TestCase):
    def test_hidden_json_contains_correct_financial_data(self):
        data, _ = run_pipeline()
        self.assertEqual(data, EXPECTED)

    def test_hidden_docx_exists_and_opens(self):
        _, doc = run_pipeline()
        self.assertGreater(len(doc.paragraphs), 0)

    def test_hidden_docx_numbers_match_pdf(self):
        _, doc = run_pipeline()
        text = doc_all_text(doc)
        observed_numbers = numeric_values(text)
        for quarter, values in EXPECTED.items():
            self.assertIn(quarter, text)
            for value in values.values():
                self.assertIn(
                    Decimal(str(value)),
                    observed_numbers,
                    f"financial value {value} missing from DOCX",
                )

    def test_hidden_docx_has_title_and_table(self):
        _, doc = run_pipeline()
        text = doc_all_text(doc).lower()
        self.assertTrue("financial report" in text or "quarterly financial report" in text)
        self.assertGreaterEqual(len(doc.tables), 1)

if __name__ == "__main__":
    unittest.main()
