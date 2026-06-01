from __future__ import annotations

from docx import Document


def build_docx(path, data: dict) -> None:
    doc = Document()
    doc.add_paragraph("Quarterly Report")
    for quarter, values in data.items():
        doc.add_paragraph(f"{quarter}: revenue={values['revenue']} expenses={values['expenses']}")
    doc.save(str(path))
